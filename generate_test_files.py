from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

SAMPLE_INPUT = Path("sample_input")


TEST_FILES = [
    {
        "filename": "ACME_NDA_2024_clear.pdf",
        "kind": "pdf",
        "expected_company": "acme",
        "expected_doc_type": "nda",
        "text": [
            "ACME Non-Disclosure Agreement",
            "This NDA governs confidential information exchanged between ACME and the recipient.",
            "The parties agree to protect trade secrets and confidential information.",
        ],
    },
    {
        "filename": "globex-master-services-agreement.docx",
        "kind": "docx",
        "expected_company": "globex",
        "expected_doc_type": "contracts",
        "text": [
            "Master Services Agreement",
            "This agreement is entered into by Globex and the customer for consulting services.",
            "The parties agree to statement of work terms, payment terms, and service levels.",
        ],
    },
    {
        "filename": "initech_invoice_1042.txt",
        "kind": "txt",
        "expected_company": "initech",
        "expected_doc_type": "financial",
        "text": [
            "Invoice #1042",
            "Vendor: Initech",
            "Amount due: 4,200 USD",
            "Payment terms: Net 30",
        ],
    },
    {
        "filename": "2024-05 balance sheet - ACME.pdf",
        "kind": "pdf",
        "expected_company": "acme",
        "expected_doc_type": "financial",
        "text": [
            "ACME Balance Sheet",
            "Assets, liabilities, revenue, expenses, and retained earnings for May 2024.",
        ],
    },
    {
        "filename": "Globex Confidentiality Agreement FINAL.docx",
        "kind": "docx",
        "expected_company": "globex",
        "expected_doc_type": "nda",
        "text": [
            "Confidentiality Agreement",
            "Globex shares confidential information under this non-disclosure agreement.",
        ],
    },
    {
        "filename": "initech_sow_phase_2.pdf",
        "kind": "pdf",
        "expected_company": "initech",
        "expected_doc_type": "contracts",
        "text": [
            "Statement of Work",
            "Initech will provide implementation services under this SOW and service agreement.",
        ],
    },
    {
        "filename": "unknown_vendor_receipt_77.txt",
        "kind": "txt",
        "expected_company": "unknown",
        "expected_doc_type": "financial",
        "text": [
            "Receipt #77",
            "Paid by card for office supplies.",
            "Vendor: Northwind Traders",
        ],
    },
    {
        "filename": "Northwind_purchase_agreement.docx",
        "kind": "docx",
        "expected_company": "unknown",
        "expected_doc_type": "contracts",
        "text": [
            "Purchase Agreement",
            "Northwind Traders agrees to sell equipment under the terms of this contract.",
        ],
    },
    {
        "filename": "scan_acme_payment_notice.png",
        "kind": "image",
        "expected_company": "acme",
        "expected_doc_type": "financial",
        "text": [
            "ACME",
            "Payment Notice",
            "Invoice balance due: 900 USD",
        ],
    },
    {
        "filename": "image_globex_nda.jpg",
        "kind": "image",
        "expected_company": "globex",
        "expected_doc_type": "nda",
        "text": [
            "Globex NDA",
            "Confidential information",
        ],
    },
    {
        "filename": "photo_initech_contract.bmp",
        "kind": "image",
        "expected_company": "initech",
        "expected_doc_type": "contracts",
        "text": [
            "Initech Contract",
            "Services Agreement",
        ],
    },
    {
        "filename": "misc_notes.txt",
        "kind": "txt",
        "expected_company": "unknown",
        "expected_doc_type": "unknown",
        "text": [
            "Meeting notes",
            "Discussed onboarding timeline and office lunch preferences.",
        ],
    },
    {
        "filename": "ACME - maybe agreement or invoice.txt",
        "kind": "txt",
        "expected_company": "acme",
        "expected_doc_type": "unknown",
        "text": [
            "ACME",
            "Agreement invoice payment contract",
            "This document has conflicting labels and little reliable context.",
        ],
    },
    {
        "filename": "untitled.pdf",
        "kind": "pdf",
        "expected_company": "unknown",
        "expected_doc_type": "unknown",
        "text": [
            "Approved.",
            "Please process this document.",
        ],
    },
    {
        "filename": "Globex_ro_contract_servicii.pdf",
        "kind": "pdf",
        "expected_company": "globex",
        "expected_doc_type": "contracts",
        "text": [
            "Contract de prestari servicii",
            "Globex si clientul convin asupra termenilor de livrare si plata.",
            "Acest acord este guvernat de legislatia aplicabila.",
        ],
    },
    {
        "filename": "Fwd Re docs from Maria.docx",
        "kind": "docx",
        "expected_company": "unknown",
        "expected_doc_type": "unknown",
        "text": [
            "Hi,",
            "Please see attached. We can discuss tomorrow.",
        ],
    },
    {
        "filename": "initech-tax-statement-2023.pdf",
        "kind": "pdf",
        "expected_company": "initech",
        "expected_doc_type": "financial",
        "text": [
            "Initech Tax Statement 2023",
            "Taxable revenue, deductions, and payment summary.",
        ],
    },
    {
        "filename": "Acme order form and NDA mixed.pdf",
        "kind": "pdf",
        "expected_company": "acme",
        "expected_doc_type": "unknown",
        "text": [
            "ACME Order Form",
            "This order form references a non-disclosure agreement and confidentiality obligations.",
            "The dominant purpose is unclear from this excerpt.",
        ],
    },
    {
        "filename": "Globex P&L Q4.csv",
        "kind": "txt",
        "expected_company": "globex",
        "expected_doc_type": "financial",
        "text": [
            "account,amount",
            "Globex revenue,100000",
            "profit and loss,25000",
        ],
    },
    {
        "filename": "x7_final_final.pdf",
        "kind": "pdf",
        "expected_company": "unknown",
        "expected_doc_type": "unknown",
        "text": [
            "Final version",
            "Reference X7",
        ],
    },
]


def _write_pdf(path: Path, lines: list[str]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    _, height = letter
    y = height - 72
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 18
    pdf.save()


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    document.save(path)


def _write_text(path: Path, lines: list[str]) -> None:
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_image(path: Path, lines: list[str]) -> None:
    image = Image.new("RGB", (1000, 500), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 42)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in lines:
        draw.text((60, y), line, fill="black", font=font)
        y += 70
    image.save(path)


def generate() -> None:
    if SAMPLE_INPUT.exists():
        shutil.rmtree(SAMPLE_INPUT)
    SAMPLE_INPUT.mkdir(parents=True)

    for item in TEST_FILES:
        path = SAMPLE_INPUT / item["filename"]
        lines = item["text"]
        if item["kind"] == "pdf":
            _write_pdf(path, lines)
        elif item["kind"] == "docx":
            _write_docx(path, lines)
        elif item["kind"] == "image":
            _write_image(path, lines)
        else:
            _write_text(path, lines)

    print(f"Generated {len(TEST_FILES)} test files in {SAMPLE_INPUT}")


if __name__ == "__main__":
    generate()
