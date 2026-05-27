from __future__ import annotations

import mimetypes
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
import pytesseract
from docx import Document
from PIL import Image

from docsort.models import AppConfig, ExtractedDocument

TEXT_SUFFIXES = {".txt", ".md", ".csv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"}


def _cap_text(text: str, config: AppConfig) -> str:
    return text[: config.text.max_chars]


def _mime_type(path: Path) -> str:
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_docx(path: Path) -> tuple[str, dict[str, Any]]:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    props = document.core_properties
    metadata = {
        "author": props.author,
        "title": props.title,
        "subject": props.subject,
        "keywords": props.keywords,
        "last_modified_by": props.last_modified_by,
    }
    return text, {key: value for key, value in metadata.items() if value}


def _ocr_image(image: Image.Image) -> str:
    return pytesseract.image_to_string(image)


def _read_pdf(path: Path, config: AppConfig) -> tuple[str, dict[str, Any], bool]:
    ocr_attempted = False
    with fitz.open(path) as document:
        metadata = dict(document.metadata or {})
        parts = [page.get_text("text") for page in document]
        text = "\n".join(parts)

        if (
            config.ocr.enabled
            and len(text.strip()) < config.ocr.min_text_chars_before_ocr
            and len(document) > 0
        ):
            ocr_attempted = True
            page = document[0]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            text = f"{text}\n{_ocr_image(image)}".strip()

    return text, metadata, ocr_attempted


def _read_image(path: Path, config: AppConfig) -> tuple[str, bool]:
    if not config.ocr.enabled:
        return "", False
    with Image.open(path) as image:
        return _ocr_image(image), True


def extract_document(path: Path, config: AppConfig) -> ExtractedDocument:
    """Extract text and metadata without raising on per-file extraction failures."""
    suffix = path.suffix.lower()
    metadata: dict[str, Any] = {}
    text = ""
    ocr_attempted = False
    extraction_error: str | None = None

    try:
        if suffix in TEXT_SUFFIXES:
            text = _read_text_file(path)
        elif suffix == ".pdf":
            text, metadata, ocr_attempted = _read_pdf(path, config)
        elif suffix == ".docx":
            text, metadata = _read_docx(path)
        elif suffix in IMAGE_SUFFIXES:
            text, ocr_attempted = _read_image(path, config)
        else:
            extraction_error = f"Unsupported file type: {suffix or '<none>'}"
    except Exception as exc:  # noqa: BLE001 - continue processing other files.
        extraction_error = f"{type(exc).__name__}: {exc}"

    return ExtractedDocument(
        source_path=path,
        filename=path.name,
        extension=suffix,
        mime_type=_mime_type(path),
        metadata=metadata,
        text=_cap_text(text, config),
        ocr_attempted=ocr_attempted,
        extraction_error=extraction_error,
    )
