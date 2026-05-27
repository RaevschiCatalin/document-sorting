from __future__ import annotations

from datetime import datetime
from itertools import count
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

_COUNTER = count(1)


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _write_pdf(path: Path, lines: list[str]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    _, height = letter
    y = height - 72
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 18
    pdf.save()


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    document.save(path)


def _write_image(path: Path, lines: list[str]) -> None:
    image = Image.new("RGB", (1100, 520), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in lines:
        draw.text((60, y), line, fill="black", font=font)
        y += 68
    image.save(path)


def create_demo_files(input_dir: Path, count_files: int = 5) -> list[Path]:
    """Create a mixed batch of demo documents in the incoming folder.

    The batch intentionally includes both clear classifications and review cases:
    - clear invoice PDF
    - clear NDA DOCX
    - ambiguous mixed-signal TXT
    - unknown-company financial CSV
    - ambiguous scanned image
    """
    input_dir.mkdir(parents=True, exist_ok=True)
    batch = next(_COUNTER)
    suffix = f"{_stamp()}_{batch:03d}"
    paths: list[Path] = []

    creators = [
        _create_invoice_pdf,
        _create_nda_docx,
        _create_ambiguous_text,
        _create_unknown_financial_csv,
        _create_ambiguous_scan_png,
    ]
    for creator in creators[:count_files]:
        paths.append(creator(input_dir, suffix))
    return paths


def create_demo_file(input_dir: Path) -> Path:
    """Backward-compatible helper that returns the first file from a batch."""
    return create_demo_files(input_dir, count_files=1)[0]


def _create_invoice_pdf(input_dir: Path, suffix: str) -> Path:
    path = input_dir / f"auto_demo_acme_invoice_{suffix}.pdf"
    _write_pdf(
        path,
        [
            "Invoice",
            "Vendor: ACME Corporation",
            "Amount due: 2,450 USD",
            "Payment terms: Net 30",
            f"Reference: {suffix}",
        ],
    )
    return path


def _create_nda_docx(input_dir: Path, suffix: str) -> Path:
    path = input_dir / f"auto_demo_globex_nda_{suffix}.docx"
    _write_docx(
        path,
            [
                "Globex Industries Non-Disclosure Agreement",
                "This NDA protects confidential information exchanged by the parties.",
                f"Reference: {suffix}",
            ],
        )
    return path


def _create_ambiguous_text(input_dir: Path, suffix: str) -> Path:
    path = input_dir / f"auto_demo_ambiguous_text_{suffix}.txt"
    path.write_text(
        "ACME Corporation\n\n"
        "Agreement invoice payment contract.\n"
        "This file intentionally mixes legal and financial labels and must be reviewed manually.\n"
        f"Reference: {suffix}\n",
        encoding="utf-8",
    )
    return path


def _create_unknown_financial_csv(input_dir: Path, suffix: str) -> Path:
    path = input_dir / f"auto_demo_unknown_receipt_{suffix}.csv"
    path.write_text(
        "vendor,description,amount\n"
        f"Northwind Traders,office supplies receipt {suffix},184.20\n",
        encoding="utf-8",
    )
    return path


def _create_ambiguous_scan_png(input_dir: Path, suffix: str) -> Path:
    path = input_dir / f"auto_demo_ambiguous_scan_{suffix}.png"
    _write_image(
        path,
        [
            "ACME",
            "Agreement invoice payment contract",
            "Ambiguous scan for manual review",
            f"Reference {suffix}",
        ],
    )
    return path
